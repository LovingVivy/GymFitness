"""Shop catalogue: public listing, admin management and Excel/CSV import (ADR-017)."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections.abc import Iterator, Mapping
from typing import Any

from fastapi import status
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.core import cache
from app.core.config import get_settings
from app.db.base import new_uuid
from app.repositories import shop as shop_repository
from app.schemas.shop import (
    AdminProduct,
    AdminProductPage,
    AdminProductSave,
    ImportRowError,
    ProductImportResult,
    ShopProduct,
)
from app.services.pt import error, raise_for_outcome

_ERRORS: dict[str, tuple[int, str]] = {
    "PRODUCT_NOT_FOUND": (status.HTTP_404_NOT_FOUND, "Không tìm thấy sản phẩm."),
    "SKU_EXISTS": (status.HTTP_409_CONFLICT, "Mã SKU này đã được dùng cho sản phẩm khác."),
    "STOCK_BELOW_RESERVED": (
        status.HTTP_409_CONFLICT,
        "Tồn kho không được thấp hơn số đang giữ cho đơn hàng.",
    ),
    "INVALID_PRODUCT": (
        status.HTTP_422_UNPROCESSABLE_CONTENT,
        "Thông tin sản phẩm không hợp lệ.",
    ),
    "VERSION_CONFLICT": (
        status.HTTP_409_CONFLICT,
        "Sản phẩm vừa được người khác sửa. Tải lại để xem bản mới nhất.",
    ),
    "INTERNAL_ERROR": (status.HTTP_500_INTERNAL_SERVER_ERROR, "Lỗi hệ thống."),
}


def slugify(text: str) -> str:
    """Vietnamese name to URL slug: "Găng tay Grip Pro" -> "gang-tay-grip-pro"."""
    folded = unicodedata.normalize("NFD", text.replace("đ", "d").replace("Đ", "D"))
    ascii_text = "".join(char for char in folded if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "-", ascii_text.lower()).strip("-") or "san-pham"


def list_products(db: Session) -> list[ShopProduct]:
    return cache.get_or_load(
        cache.CATALOG,
        "products",
        get_settings().cache_catalog_ttl_seconds,
        lambda: [ShopProduct(**row) for row in shop_repository.list_public(db)],
        list[ShopProduct],
    )


def _to_admin_product(row: Mapping[str, Any]) -> AdminProduct:
    return AdminProduct(**{key: row.get(key) for key in AdminProduct.model_fields})


def list_admin_products(
    db: Session,
    *,
    search: str | None,
    activity_type: str | None,
    product_status: str | None,
    page: int,
    page_size: int,
    sort: str | None = None,
) -> AdminProductPage:
    rows = shop_repository.list_admin(
        db,
        search=search,
        activity_type=activity_type,
        status=product_status,
        sort=sort,
        limit=page_size,
        offset=(page - 1) * page_size,
    )
    total = int(rows[0]["total_count"]) if rows else 0
    return AdminProductPage(items=[_to_admin_product(row) for row in rows], total=total)


def _find_product(db: Session, product_id: str) -> AdminProduct | None:
    rows = shop_repository.list_admin(
        db, product_id=product_id, search=None, activity_type=None, status=None, limit=1, offset=0
    )
    return _to_admin_product(rows[0]) if rows else None


def save_product(
    db: Session, actor_id: str, payload: AdminProductSave, product_id: str | None = None
) -> AdminProduct:
    is_new = product_id is None
    target_id = product_id or new_uuid()

    sku = payload.sku
    if sku is None:
        if is_new:
            sku = f"SP-{target_id[:8].upper()}"
        else:
            current = _find_product(db, target_id)
            if current is None:
                raise error(
                    status.HTTP_404_NOT_FOUND, "PRODUCT_NOT_FOUND", "Không tìm thấy sản phẩm."
                )
            sku = current.sku

    outcome = shop_repository.save_product(
        db,
        product_id=target_id,
        variant_id=new_uuid(),
        is_new=is_new,
        # Hậu tố ngẫu nhiên để hai sản phẩm trùng tên không đụng slug.
        slug=f"{slugify(payload.name)[:160]}-{target_id[:8]}",
        name=payload.name,
        sku=sku,
        category_name=payload.category_name,
        activity_type=payload.activity_type,
        description=payload.description,
        price=payload.price,
        stock=payload.stock_quantity,
        status=payload.status,
        badge=payload.badge,
        art=payload.art,
        expected_version=None if is_new else payload.row_version,
        actor_id=actor_id,
    )
    if not outcome.succeeded:
        raise_for_outcome(outcome, _ERRORS, "Không thể lưu sản phẩm.")
    cache.invalidate(cache.CATALOG)

    product = _find_product(db, target_id)
    if product is None:
        raise error(
            status.HTTP_500_INTERNAL_SERVER_ERROR, "INTERNAL_ERROR", "Không đọc lại được sản phẩm."
        )
    return product


def delete_product(db: Session, actor_id: str, product_id: str) -> None:
    outcome = shop_repository.delete_product(db, product_id=product_id, actor_id=actor_id)
    if not outcome.succeeded:
        raise_for_outcome(outcome, _ERRORS, "Không thể xoá sản phẩm.")
    cache.invalidate(cache.CATALOG)


# --------------------------------------------------------------------------
# Import from Excel / CSV
# --------------------------------------------------------------------------

MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 2_000

# (khoá trong AdminProductSave, tiêu đề cột trong file mẫu)
IMPORT_COLUMNS: tuple[tuple[str, str], ...] = (
    ("sku", "SKU"),
    ("name", "Tên sản phẩm"),
    ("category_name", "Danh mục"),
    ("activity_type", "Nhóm (GYM/YOGA)"),
    ("price", "Giá"),
    ("stock_quantity", "Tồn kho"),
    ("status", "Trạng thái (ACTIVE/DRAFT/INACTIVE)"),
    ("description", "Mô tả"),
    ("badge", "Nhãn"),
)
_REQUIRED_COLUMNS = frozenset(
    {"sku", "name", "category_name", "activity_type", "price", "stock_quantity"}
)

_SAMPLE_ROWS = (
    (
        "GYM-ROPE-01",
        "Dây nhảy tốc độ",
        "Phụ kiện",
        "GYM",
        189000,
        40,
        "ACTIVE",
        "Cáp thép bọc PVC, tay cầm chống trượt",
        "Mới",
    ),
    (
        "YOGA-WHEEL-01",
        "Bánh xe Yoga",
        "Dụng cụ",
        "YOGA",
        359000,
        15,
        "DRAFT",
        "Hỗ trợ giãn lưng và mở vai",
        "",
    ),
)


def _fold(text: str) -> str:
    folded = unicodedata.normalize("NFD", text.replace("đ", "d").replace("Đ", "D"))
    plain = "".join(char for char in folded if not unicodedata.combining(char)).lower()
    return re.sub(r"[^a-z0-9]+", "", plain)


# Chấp nhận cả khoá tiếng Anh lẫn tiêu đề tiếng Việt (có hoặc không dấu, bỏ phần ngoặc).
_HEADER_KEYS = {
    **{_fold(key): key for key, _ in IMPORT_COLUMNS},
    **{_fold(label.split("(")[0]): key for key, label in IMPORT_COLUMNS},
}


def _header_key(value: Any) -> str | None:
    """Column key of a header cell such as "Nhóm (GYM/YOGA)", "Ton kho" or "price"."""
    return _HEADER_KEYS.get(_fold(str(value or "").split("(")[0]))


def import_template(fmt: str) -> tuple[str, str, bytes]:
    """Empty template with two sample rows: (file name, content type, bytes)."""
    headers = [label for _, label in IMPORT_COLUMNS]
    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(headers)
        writer.writerows(_SAMPLE_ROWS)
        content = ("﻿" + buffer.getvalue()).encode("utf-8")
        return "mau-nhap-san-pham.csv", "text/csv; charset=utf-8", content

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "San pham"
    sheet.append(headers)
    for row in _SAMPLE_ROWS:
        sheet.append(list(row))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="111111")
    for letter, width in zip("ABCDEFGHI", (16, 28, 16, 16, 12, 10, 30, 40, 12), strict=True):
        sheet.column_dimensions[letter].width = width
    buffer = io.BytesIO()
    workbook.save(buffer)
    return (
        "mau-nhap-san-pham.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        buffer.getvalue(),
    )


def _rows_from_file(filename: str, content: bytes) -> Iterator[tuple[int, list[Any]]]:
    """Yield (row number in the file, cell values), header row included."""
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise error(
                status.HTTP_422_UNPROCESSABLE_CONTENT,
                "INVALID_FILE",
                "File CSV phải được lưu bằng mã UTF-8.",
            ) from exc
        for number, row in enumerate(csv.reader(io.StringIO(text)), start=1):
            yield number, list(row)
        return
    if name.endswith(".xlsx"):
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - openpyxl ném nhiều loại lỗi khác nhau
            raise error(
                status.HTTP_422_UNPROCESSABLE_CONTENT,
                "INVALID_FILE",
                "Không đọc được file Excel (.xlsx).",
            ) from exc
        try:
            sheet = workbook.worksheets[0]
            for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                yield number, list(row)
        finally:
            workbook.close()
        return
    raise error(
        status.HTTP_422_UNPROCESSABLE_CONTENT,
        "UNSUPPORTED_FILE",
        "Chỉ nhận file .xlsx hoặc .csv.",
    )


def _cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _validation_message(exc: ValidationError) -> str:
    labels = dict(IMPORT_COLUMNS)
    parts = []
    for item in exc.errors(include_url=False):
        field = str(item["loc"][0]) if item.get("loc") else ""
        message = str(item.get("msg", "không hợp lệ")).removeprefix("Value error, ")
        parts.append(f"{labels.get(field, field)}: {message}")
    return "; ".join(parts)


def _normalise(record: dict[str, Any]) -> dict[str, Any]:
    for key in ("activity_type", "status"):
        if isinstance(record.get(key), str):
            record[key] = record[key].upper()
    for key in ("sku", "name", "category_name", "description", "badge"):
        if key in record and not isinstance(record[key], str):
            record[key] = str(record[key])
    return record


def import_products(
    db: Session, actor_id: str, filename: str, content: bytes
) -> ProductImportResult:
    """Create or update products from a spreadsheet, keyed by SKU.

    Each row is saved on its own: a bad row is reported back and the others
    still go in, so a long file does not have to be fixed in one go.
    """
    if len(content) > MAX_IMPORT_BYTES:
        raise error(status.HTTP_413_CONTENT_TOO_LARGE, "FILE_TOO_LARGE", "File nhập tối đa 2MB.")

    rows = _rows_from_file(filename, content)
    header_keys: list[str | None] = []
    for _, values in rows:
        if any(_cell(value) is not None for value in values):
            header_keys = [_header_key(value) for value in values]
            break
    missing = _REQUIRED_COLUMNS - {key for key in header_keys if key}
    if missing:
        labels = dict(IMPORT_COLUMNS)
        raise error(
            status.HTTP_422_UNPROCESSABLE_CONTENT,
            "MISSING_COLUMNS",
            "Thiếu cột: " + ", ".join(labels[key] for key in sorted(missing)) + ".",
        )

    created = updated = processed = 0
    errors: list[ImportRowError] = []
    for number, values in rows:
        record = {
            key: _cell(value)
            for key, value in zip(header_keys, values, strict=False)
            if key and _cell(value) is not None
        }
        if not record:
            continue
        processed += 1
        if processed > MAX_IMPORT_ROWS:
            errors.append(
                ImportRowError(row=number, message=f"Mỗi lần nhập tối đa {MAX_IMPORT_ROWS} dòng.")
            )
            break
        try:
            payload = AdminProductSave(**_normalise(record))
        except ValidationError as exc:
            errors.append(ImportRowError(row=number, message=_validation_message(exc)))
            continue
        if payload.sku is None:
            errors.append(ImportRowError(row=number, message="SKU: không được để trống"))
            continue

        existing = shop_repository.find_by_sku(db, payload.sku)
        try:
            save_product(db, actor_id, payload, existing["product_id"] if existing else None)
        except Exception as exc:  # noqa: BLE001 - lỗi của một dòng không dừng cả file
            detail = getattr(exc, "detail", None)
            message = (
                detail["error"]["message"]
                if isinstance(detail, dict) and "error" in detail
                else "Không lưu được dòng này."
            )
            errors.append(ImportRowError(row=number, message=message))
            continue
        if existing:
            updated += 1
        else:
            created += 1

    return ProductImportResult(created=created, updated=updated, errors=errors)
