"""Integration tests for the platform features: RBAC (Manager), sorting, optimistic
locking, audit log, exports through the job queue, invoices, spreadsheet import,
password reset, two-step verification and catalogue caching."""

from __future__ import annotations

import csv
import io
import re
import uuid
from collections.abc import Iterator

import pytest
from openpyxl import Workbook, load_workbook

from app.core import cache, totp
from app.core.config import get_settings
from tests.conftest import client, mail_text, purge_user

ADMIN_URL = "/api/v1/admin"
AUTH_URL = "/api/v1/auth"


def _promote(db, session: dict, role: str) -> dict:
    with db.cursor() as cursor:
        cursor.execute("UPDATE users SET role = %s WHERE id = %s", (role, session["user"]["id"]))
    return session


@pytest.fixture
def admin(session, db) -> dict:
    return _promote(db, session, "ADMIN")


@pytest.fixture
def manager(account, db) -> Iterator[dict]:
    """A second account promoted to MANAGER (its own email, purged afterwards)."""
    email = f"pytest.manager.{uuid.uuid4().hex[:10]}@example.com"
    response = client.post(
        f"{AUTH_URL}/register",
        json={"full_name": "Quản lý Pytest", "email": email, "password": "password123"},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    with db.cursor() as cursor:
        cursor.execute("UPDATE users SET role = 'MANAGER' WHERE id = %s", (body["user"]["id"],))
    yield {
        "user": body["user"],
        "email": email,
        "headers": {"Authorization": f"Bearer {body['access_token']}"},
    }
    purge_user(db, body["user"]["id"])


@pytest.fixture
def created_users(db) -> Iterator[list[str]]:
    ids: list[str] = []
    yield ids
    for user_id in ids:
        purge_user(db, user_id)


@pytest.fixture
def cleanup_catalog(db) -> Iterator[None]:
    yield
    with db.cursor() as cursor:
        cursor.execute(
            "DELETE FROM audit_logs WHERE entity_type = 'membership_plan' AND entity_id IN "
            "(SELECT id FROM membership_plans WHERE code LIKE 'PYTEST_%%')"
        )
        cursor.execute("DELETE FROM membership_plans WHERE code LIKE 'PYTEST_%%'")
        cursor.execute(
            "SELECT p.id FROM products p JOIN product_variants v ON v.product_id = p.id "
            "WHERE v.sku LIKE 'PYIMP-%%' OR v.sku LIKE 'DEL-%%PYIMP-%%' OR p.name LIKE 'Pytest %%'"
        )
        for (product_id,) in cursor.fetchall():
            cursor.execute(
                "DELETE FROM inventory_movements WHERE variant_id IN "
                "(SELECT id FROM product_variants WHERE product_id = %s)",
                (product_id,),
            )
            cursor.execute("DELETE FROM product_variants WHERE product_id = %s", (product_id,))
            cursor.execute(
                "DELETE FROM audit_logs WHERE entity_type = 'product' AND entity_id = %s",
                (product_id,),
            )
            cursor.execute("DELETE FROM products WHERE id = %s", (product_id,))
        cursor.execute("DELETE FROM product_categories WHERE name LIKE 'Pytest %%'")


def _plan(code: str = "GYM_MONTHLY") -> dict:
    return next(item for item in client.get("/api/v1/plans").json() if item["code"] == code)


def _subscribe(session: dict, pay: bool = True, pt_sessions: int = 0) -> dict:
    created = client.post(
        "/api/v1/subscriptions",
        json={"plan_id": _plan("GYM_YEARLY")["id"], "pt_sessions": pt_sessions},
        headers=session["headers"],
    )
    assert created.status_code == 201, created.text
    payment = created.json()["payment"]
    if pay:
        paid = client.post(
            f"/api/v1/payments/{payment['id']}/mark-paid", headers=session["headers"]
        )
        assert paid.status_code == 200, paid.text
    return payment


# --------------------------------------------------------------------------
# RBAC: Manager
# --------------------------------------------------------------------------


def test_login_and_profile_expose_role_permissions(manager, session) -> None:
    profile = client.get("/api/v1/profile", headers=manager["headers"]).json()
    assert profile["role"] == "MANAGER"
    assert "payments:confirm" in profile["permissions"]
    assert "audit:view" not in profile["permissions"]
    member = client.get("/api/v1/profile", headers=session["headers"]).json()
    assert member["permissions"] == []
    assert member["two_factor_enabled"] is False


def test_manager_can_run_daily_operations(manager, created_users) -> None:
    headers = manager["headers"]
    for path in (
        "/overview",
        "/users",
        "/plans",
        "/checkins",
        "/revenue",
        "/payments",
        "/products",
    ):
        assert client.get(f"{ADMIN_URL}{path}", headers=headers).status_code == 200, path

    member = client.post(
        f"{ADMIN_URL}/users",
        json={
            "full_name": "Hội viên do quản lý tạo",
            "email": f"pytest.mgr.{uuid.uuid4().hex[:8]}@example.com",
            "password": "password123",
        },
        headers=headers,
    )
    assert member.status_code == 201, member.text
    created_users.append(member.json()["id"])
    updated = client.patch(
        f"{ADMIN_URL}/users/{member.json()['id']}", json={"full_name": "Đã sửa"}, headers=headers
    )
    assert updated.status_code == 200, updated.text


def test_manager_cannot_touch_staff_plans_voids_or_audit(manager, admin, created_users) -> None:
    headers = manager["headers"]
    trainer = client.post(
        f"{ADMIN_URL}/users",
        json={
            "full_name": "PT",
            "email": f"pytest.pt.{uuid.uuid4().hex[:8]}@example.com",
            "password": "password123",
            "role": "PT",
            "specialties": ["GYM"],
        },
        headers=headers,
    )
    assert trainer.status_code == 403
    assert trainer.json()["error"]["code"] == "FORBIDDEN"

    # Tài khoản quản trị: quản lý không sửa, không đặt lại mật khẩu, không xoá được.
    admin_id = admin["user"]["id"]
    assert (
        client.patch(
            f"{ADMIN_URL}/users/{admin_id}", json={"full_name": "x"}, headers=headers
        ).status_code
        == 403
    )
    assert (
        client.post(
            f"{ADMIN_URL}/users/{admin_id}/password",
            json={"password": "password999"},
            headers=headers,
        ).status_code
        == 403
    )
    assert client.delete(f"{ADMIN_URL}/users/{admin_id}", headers=headers).status_code == 403

    plan_body = {
        "code": "PYTEST_MGR",
        "name": "Không được tạo",
        "activity_type": "GYM",
        "duration_days": 7,
        "price": 1,
    }
    assert client.post(f"{ADMIN_URL}/plans", json=plan_body, headers=headers).status_code == 403
    assert (
        client.post(f"{ADMIN_URL}/payments/{uuid.uuid4()}/void", headers=headers).status_code == 403
    )
    assert client.get(f"{ADMIN_URL}/audit-logs", headers=headers).status_code == 403


def test_admin_can_create_a_manager(admin, created_users) -> None:
    response = client.post(
        f"{ADMIN_URL}/users",
        json={
            "full_name": "Quản lý mới",
            "email": f"pytest.newmgr.{uuid.uuid4().hex[:8]}@example.com",
            "password": "password123",
            "role": "MANAGER",
        },
        headers=admin["headers"],
    )
    assert response.status_code == 201, response.text
    created_users.append(response.json()["id"])
    assert response.json()["role"] == "MANAGER"


def test_manager_confirms_payment_at_the_desk(manager, session) -> None:
    payment = _subscribe(session, pay=False)
    response = client.post(
        f"{ADMIN_URL}/payments/{payment['id']}/confirm",
        json={"method": "CASH"},
        headers=manager["headers"],
    )
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "PAID"


# --------------------------------------------------------------------------
# Sorting
# --------------------------------------------------------------------------


def _is_sorted(values: list, reverse: bool = False) -> bool:
    present = [value for value in values if value is not None]
    return present == sorted(present, reverse=reverse)


@pytest.mark.parametrize(
    ("path", "sort", "field", "reverse"),
    [
        ("/users", "email", "email", False),
        ("/users", "-email", "email", True),
        ("/payments", "-amount", "amount", True),
        ("/payments", "amount", "amount", False),
        ("/products", "price", "price", False),
        ("/products", "-stock_quantity", "stock_quantity", True),
        ("/checkins", "check_in_at", "check_in_at", False),
    ],
)
def test_admin_lists_sort_by_the_requested_column(admin, path, sort, field, reverse) -> None:
    response = client.get(
        f"{ADMIN_URL}{path}", params={"sort": sort, "page_size": 50}, headers=admin["headers"]
    )
    assert response.status_code == 200, response.text
    assert _is_sorted([item[field] for item in response.json()["items"]], reverse)


def test_unknown_sort_key_is_rejected(admin) -> None:
    response = client.get(
        f"{ADMIN_URL}/users", params={"sort": "password_hash"}, headers=admin["headers"]
    )
    assert response.status_code == 422


def test_sorted_pages_do_not_overlap(admin) -> None:
    params = {"sort": "full_name", "page_size": 5}
    first = client.get(f"{ADMIN_URL}/users", params={**params, "page": 1}, headers=admin["headers"])
    second = client.get(
        f"{ADMIN_URL}/users", params={**params, "page": 2}, headers=admin["headers"]
    )
    ids_first = {item["id"] for item in first.json()["items"]}
    ids_second = {item["id"] for item in second.json()["items"]}
    assert ids_first and not ids_first & ids_second


# --------------------------------------------------------------------------
# Optimistic locking and audit log
# --------------------------------------------------------------------------


def test_plan_update_with_stale_version_is_rejected(admin, cleanup_catalog) -> None:
    body = {
        "code": f"PYTEST_V{uuid.uuid4().hex[:6].upper()}",
        "name": "Gói phiên bản",
        "activity_type": "GYM",
        "duration_days": 10,
        "price": 100000,
    }
    created = client.post(f"{ADMIN_URL}/plans", json=body, headers=admin["headers"]).json()
    assert created["row_version"] == 1

    first = client.put(
        f"{ADMIN_URL}/plans/{created['id']}",
        json={**body, "price": 120000, "row_version": 1},
        headers=admin["headers"],
    )
    assert first.status_code == 200, first.text
    assert first.json()["row_version"] == 2

    stale = client.put(
        f"{ADMIN_URL}/plans/{created['id']}",
        json={**body, "price": 90000, "row_version": 1},
        headers=admin["headers"],
    )
    assert stale.status_code == 409
    assert stale.json()["error"]["code"] == "VERSION_CONFLICT"

    history = client.get(
        f"{ADMIN_URL}/audit-logs",
        params={"entity_type": "membership_plan", "entity_id": created["id"]},
        headers=admin["headers"],
    ).json()
    assert [item["action"] for item in history["items"]] == ["PLAN_UPDATE", "PLAN_CREATE"]
    update = history["items"][0]
    assert update["actor_user_id"] == admin["user"]["id"]
    assert update["before_data"]["row_version"] == 1
    assert float(update["after_data"]["price"]) == 120000


def test_product_update_with_stale_version_is_rejected(admin, cleanup_catalog) -> None:
    body = {
        "name": f"Pytest Bình nước {uuid.uuid4().hex[:6]}",
        "category_name": "Pytest Phụ kiện",
        "activity_type": "GYM",
        "price": 99000,
        "stock_quantity": 5,
    }
    created = client.post(f"{ADMIN_URL}/products", json=body, headers=admin["headers"]).json()
    ok = client.put(
        f"{ADMIN_URL}/products/{created['id']}",
        json={**body, "row_version": created["row_version"]},
        headers=admin["headers"],
    )
    assert ok.status_code == 200, ok.text
    stale = client.put(
        f"{ADMIN_URL}/products/{created['id']}",
        json={**body, "row_version": created["row_version"]},
        headers=admin["headers"],
    )
    assert stale.status_code == 409
    assert stale.json()["error"]["code"] == "VERSION_CONFLICT"


def test_audit_log_filters_and_pagination(admin) -> None:
    page = client.get(
        f"{ADMIN_URL}/audit-logs", params={"page_size": 2}, headers=admin["headers"]
    ).json()
    assert len(page["items"]) <= 2
    assert page["total"] >= len(page["items"])
    empty = client.get(
        f"{ADMIN_URL}/audit-logs",
        params={"action": "NO_SUCH_ACTION"},
        headers=admin["headers"],
    ).json()
    assert empty == {"items": [], "total": 0}


# --------------------------------------------------------------------------
# Exports through the job queue
# --------------------------------------------------------------------------


def _export(headers: dict, **body) -> dict:
    response = client.post(f"{ADMIN_URL}/exports", json=body, headers=headers)
    assert response.status_code == 202, response.text
    return response.json()


@pytest.mark.parametrize(
    ("report", "fmt", "filters"),
    [
        ("payments", "xlsx", {"status": "PAID", "sort": "-amount"}),
        ("users", "csv", {"role": "USER"}),
        ("checkins", "pdf", {}),
        ("revenue", "xlsx", {"period": "QUARTER", "count": 4}),
        ("member_growth", "csv", {"period": "MONTH"}),
        ("products", "pdf", {"activity_type": "YOGA"}),
        ("audit_logs", "csv", {}),
    ],
)
def test_export_job_produces_a_downloadable_file(admin, report, fmt, filters) -> None:
    job = _export(admin["headers"], report=report, format=fmt, filters=filters)
    # Test chạy hàng đợi ở chế độ inline nên tác vụ đã xong khi request trả về.
    status = client.get(f"{ADMIN_URL}/jobs/{job['id']}", headers=admin["headers"]).json()
    assert status["status"] == "SUCCESS", status
    assert status["file_name"].endswith(f".{fmt}")

    download = client.get(status["download_url"], headers=admin["headers"])
    assert download.status_code == 200
    if fmt == "xlsx":
        sheet = load_workbook(io.BytesIO(download.content)).active
        assert sheet["A1"].value
    elif fmt == "csv":
        assert download.content.startswith("﻿".encode())
    else:
        assert download.content.startswith(b"%PDF-")

    listed = client.get(f"{ADMIN_URL}/jobs", headers=admin["headers"]).json()
    assert job["id"] in {item["id"] for item in listed}


def test_payment_export_rows_match_the_filtered_list(admin) -> None:
    job = _export(admin["headers"], report="payments", format="csv", filters={"status": "PAID"})
    status = client.get(f"{ADMIN_URL}/jobs/{job['id']}", headers=admin["headers"]).json()
    rows = list(
        csv.reader(
            io.StringIO(
                client.get(status["download_url"], headers=admin["headers"]).content.decode(
                    "utf-8-sig"
                )
            )
        )
    )
    total = client.get(
        f"{ADMIN_URL}/payments", params={"status": "PAID"}, headers=admin["headers"]
    ).json()["total"]
    assert rows[0][0] == "Mã giao dịch"
    assert len(rows) - 1 == total


def test_exports_are_private_and_permission_checked(admin, manager, db) -> None:
    job = _export(admin["headers"], report="users", format="csv")
    status = client.get(f"{ADMIN_URL}/jobs/{job['id']}", headers=admin["headers"]).json()
    # Người khác (kể cả quản lý) không thấy tác vụ cũng không tải được file.
    assert (
        client.get(f"{ADMIN_URL}/jobs/{job['id']}", headers=manager["headers"]).status_code == 404
    )
    assert client.get(status["download_url"], headers=manager["headers"]).status_code == 403

    denied = client.post(
        f"{ADMIN_URL}/exports", json={"report": "audit_logs"}, headers=manager["headers"]
    )
    assert denied.status_code == 403
    _promote(db, admin, "USER")
    assert (
        client.post(
            f"{ADMIN_URL}/exports", json={"report": "users"}, headers=admin["headers"]
        ).status_code
        == 403
    )


def test_export_rejects_invalid_filters(admin) -> None:
    response = client.post(
        f"{ADMIN_URL}/exports",
        json={"report": "payments", "filters": {"status": "MAYBE"}},
        headers=admin["headers"],
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "INVALID_FILTERS"


# --------------------------------------------------------------------------
# Invoices and receipt email
# --------------------------------------------------------------------------


def test_paid_membership_gets_invoice_pdf_and_receipt_email(
    session, admin, account, outbox
) -> None:
    payment = _subscribe(session, pt_sessions=2)

    invoice = client.get(f"/api/v1/payments/{payment['id']}/invoice", headers=session["headers"])
    assert invoice.status_code == 200, invoice.text
    assert invoice.headers["content-type"] == "application/pdf"
    assert invoice.content.startswith(b"%PDF-")
    number = re.search(r'filename="(HD\d{6}-\d{5})\.pdf"', invoice.headers["content-disposition"])
    assert number

    # Gọi lại trả đúng hoá đơn cũ, không lập số mới.
    again = client.get(f"/api/v1/payments/{payment['id']}/invoice", headers=session["headers"])
    assert number.group(1) in again.headers["content-disposition"]

    receipts = [mail for mail in outbox if mail["To"] == account["email"]]
    assert len(receipts) == 1
    assert number.group(1) in receipts[0]["Subject"]
    attachments = list(receipts[0].iter_attachments())
    assert attachments[0].get_filename() == f"{number.group(1)}.pdf"
    assert "Buổi tập với huấn luyện viên" in mail_text(receipts[0])

    staff_copy = client.get(f"/api/v1/payments/{payment['id']}/invoice", headers=admin["headers"])
    assert staff_copy.status_code == 200


def test_invoice_is_private_and_needs_a_paid_payment(session, manager) -> None:
    pending = _subscribe(session, pay=False)
    response = client.get(f"/api/v1/payments/{pending['id']}/invoice", headers=session["headers"])
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "PAYMENT_NOT_PAID"

    stranger = client.post(
        f"{AUTH_URL}/register",
        json={
            "full_name": "Người lạ",
            "email": f"pytest.stranger.{uuid.uuid4().hex[:8]}@example.com",
            "password": "password123",
        },
    ).json()
    try:
        response = client.get(
            f"/api/v1/payments/{pending['id']}/invoice",
            headers={"Authorization": f"Bearer {stranger['access_token']}"},
        )
        assert response.status_code == 404
    finally:
        from tests.conftest import connect

        connection = connect()
        try:
            purge_user(connection, stranger["user"]["id"])
        finally:
            connection.close()


# --------------------------------------------------------------------------
# Spreadsheet import
# --------------------------------------------------------------------------


def _upload(headers: dict, name: str, content: bytes):
    return client.post(
        f"{ADMIN_URL}/products/import",
        files={"file": (name, content, "application/octet-stream")},
        headers=headers,
    )


def test_csv_import_creates_then_updates_by_sku(manager, cleanup_catalog) -> None:
    sku = f"PYIMP-{uuid.uuid4().hex[:6].upper()}"
    content = (
        "SKU,Tên sản phẩm,Danh mục,Nhóm (GYM/YOGA),Giá,Tồn kho,Trạng thái\n"
        f"{sku},Pytest Dây kháng lực,Pytest Phụ kiện,gym,150000,12,ACTIVE\n"
        "PYIMP-BAD,Pytest Hỏng,Pytest Phụ kiện,BOXING,-5,1,ACTIVE\n"
    ).encode()
    result = _upload(manager["headers"], "san-pham.csv", content)
    assert result.status_code == 200, result.text
    body = result.json()
    assert (body["created"], body["updated"]) == (1, 0)
    assert [error["row"] for error in body["errors"]] == [3]
    assert "Nhóm" in body["errors"][0]["message"] and "Giá" in body["errors"][0]["message"]

    updated = _upload(
        manager["headers"],
        "san-pham.csv",
        f"sku,name,category_name,activity_type,price,stock_quantity\n"
        f"{sku},Pytest Dây kháng lực v2,Pytest Phụ kiện,GYM,175000,20\n".encode(),
    ).json()
    assert (updated["created"], updated["updated"], updated["errors"]) == (0, 1, [])
    listed = client.get(
        f"{ADMIN_URL}/products", params={"search": sku}, headers=manager["headers"]
    ).json()["items"]
    assert [(item["price"], item["stock_quantity"]) for item in listed] == [(175000, 20)]


def test_xlsx_import_and_template(admin, cleanup_catalog) -> None:
    template = client.get(
        f"{ADMIN_URL}/products/import-template", params={"format": "xlsx"}, headers=admin["headers"]
    )
    assert template.status_code == 200
    assert "attachment" in template.headers["content-disposition"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.append([cell.value for cell in load_workbook(io.BytesIO(template.content)).active[1]])
    sku = f"PYIMP-{uuid.uuid4().hex[:6].upper()}"
    sheet.append([sku, "Pytest Thảm tập", "Pytest Dụng cụ", "YOGA", 450000, 8, "DRAFT", "", ""])
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = _upload(admin["headers"], "san-pham.xlsx", buffer.getvalue()).json()
    assert (result["created"], result["errors"]) == (1, [])


def test_import_rejects_missing_columns_and_other_files(admin) -> None:
    missing = _upload(admin["headers"], "x.csv", b"SKU,Ten san pham\nA,B\n")
    assert missing.status_code == 422
    assert missing.json()["error"]["code"] == "MISSING_COLUMNS"
    wrong = _upload(admin["headers"], "x.pdf", b"%PDF-1.4")
    assert wrong.json()["error"]["code"] == "UNSUPPORTED_FILE"


# --------------------------------------------------------------------------
# Password reset
# --------------------------------------------------------------------------


def _reset_token(message) -> str:
    match = re.search(r"/reset-password\?token=([A-Za-z0-9_-]+)", mail_text(message))
    assert match, mail_text(message)
    return match.group(1)


def test_forgot_and_reset_password(session, account, outbox) -> None:
    response = client.post(f"{AUTH_URL}/forgot-password", json={"email": account["email"]})
    assert response.status_code == 202
    assert len(outbox) == 1 and outbox[0]["To"] == account["email"]
    token = _reset_token(outbox[0])

    reset = client.post(
        f"{AUTH_URL}/reset-password", json={"token": token, "password": "mat-khau-moi-123"}
    )
    assert reset.status_code == 200, reset.text

    old = client.post(
        f"{AUTH_URL}/login", json={"email": account["email"], "password": account["password"]}
    )
    assert old.status_code == 401
    new = client.post(
        f"{AUTH_URL}/login", json={"email": account["email"], "password": "mat-khau-moi-123"}
    )
    assert new.status_code == 200

    reused = client.post(
        f"{AUTH_URL}/reset-password", json={"token": token, "password": "lan-thu-hai-123"}
    )
    assert reused.status_code == 400
    assert reused.json()["error"]["code"] == "TOKEN_INVALID"


def test_forgot_password_does_not_reveal_unknown_emails(outbox) -> None:
    response = client.post(
        f"{AUTH_URL}/forgot-password", json={"email": f"nobody.{uuid.uuid4().hex}@example.com"}
    )
    assert response.status_code == 202
    assert outbox == []


def test_expired_and_superseded_reset_links_are_refused(session, account, db, outbox) -> None:
    client.post(f"{AUTH_URL}/forgot-password", json={"email": account["email"]})
    client.post(f"{AUTH_URL}/forgot-password", json={"email": account["email"]})
    first, second = (_reset_token(message) for message in outbox)

    superseded = client.post(
        f"{AUTH_URL}/reset-password", json={"token": first, "password": "password456"}
    )
    assert superseded.json()["error"]["code"] == "TOKEN_INVALID"

    with db.cursor() as cursor:
        cursor.execute(
            # MySQL gán từ trái sang phải: expires_at dùng created_at đã lùi 1 giờ.
            "UPDATE password_reset_tokens SET created_at = created_at - INTERVAL 1 HOUR, "
            "expires_at = created_at + INTERVAL 1 SECOND WHERE user_id = %s AND used_at IS NULL",
            (session["user"]["id"],),
        )
    expired = client.post(
        f"{AUTH_URL}/reset-password", json={"token": second, "password": "password456"}
    )
    assert expired.json()["error"]["code"] == "TOKEN_EXPIRED"


# --------------------------------------------------------------------------
# Two-step verification
# --------------------------------------------------------------------------


def _enable_two_factor(session: dict) -> str:
    setup = client.post(f"{AUTH_URL}/2fa/setup", headers=session["headers"])
    assert setup.status_code == 200, setup.text
    body = setup.json()
    assert body["qr_svg"].startswith("<svg")
    wrong = client.post(
        f"{AUTH_URL}/2fa/enable",
        json={"setup_token": body["setup_token"], "code": "000000"},
        headers=session["headers"],
    )
    if wrong.status_code == 200:  # pragma: no cover - 1/10⁶: "000000" là mã đúng
        return body["secret"]
    assert wrong.json()["error"]["code"] == "INVALID_2FA_CODE"
    enabled = client.post(
        f"{AUTH_URL}/2fa/enable",
        json={"setup_token": body["setup_token"], "code": totp.code_at(body["secret"])},
        headers=session["headers"],
    )
    assert enabled.status_code == 200, enabled.text
    return body["secret"]


def test_two_step_login(session, account) -> None:
    secret = _enable_two_factor(session)
    assert client.get("/api/v1/profile", headers=session["headers"]).json()["two_factor_enabled"]

    first = client.post(
        f"{AUTH_URL}/login", json={"email": account["email"], "password": account["password"]}
    )
    assert first.status_code == 200
    challenge = first.json()
    assert challenge["two_factor_required"] is True
    assert "access_token" not in challenge

    # Token thử thách không dùng được như phiên đăng nhập.
    assert (
        client.get(
            "/api/v1/profile", headers={"Authorization": f"Bearer {challenge['challenge_token']}"}
        ).status_code
        == 401
    )

    wrong = client.post(
        f"{AUTH_URL}/2fa/verify",
        json={"challenge_token": challenge["challenge_token"], "code": "000000"},
    )
    if wrong.status_code != 200:
        assert wrong.json()["error"]["code"] == "INVALID_2FA_CODE"

    verified = client.post(
        f"{AUTH_URL}/2fa/verify",
        json={"challenge_token": challenge["challenge_token"], "code": totp.code_at(secret)},
    )
    assert verified.status_code == 200, verified.text
    assert verified.json()["user"]["two_factor_enabled"] is True

    # Tắt cần đúng cả mật khẩu và mã.
    headers = {"Authorization": f"Bearer {verified.json()['access_token']}"}
    bad = client.post(
        f"{AUTH_URL}/2fa/disable",
        json={"password": "sai-mat-khau", "code": totp.code_at(secret)},
        headers=headers,
    )
    assert bad.status_code == 401
    disabled = client.post(
        f"{AUTH_URL}/2fa/disable",
        json={"password": account["password"], "code": totp.code_at(secret)},
        headers=headers,
    )
    assert disabled.status_code == 200
    plain = client.post(
        f"{AUTH_URL}/login", json={"email": account["email"], "password": account["password"]}
    )
    assert "access_token" in plain.json()


def test_two_factor_setup_rules(session) -> None:
    _enable_two_factor(session)
    again = client.post(f"{AUTH_URL}/2fa/setup", headers=session["headers"])
    assert again.json()["error"]["code"] == "TWO_FACTOR_ALREADY_ENABLED"
    expired = client.post(
        f"{AUTH_URL}/2fa/verify", json={"challenge_token": "x" * 40, "code": "123456"}
    )
    assert expired.json()["error"]["code"] == "CHALLENGE_EXPIRED"


# --------------------------------------------------------------------------
# Catalogue cache
# --------------------------------------------------------------------------


def test_cached_plan_list_is_invalidated_by_admin_changes(
    admin, monkeypatch, cleanup_catalog
) -> None:
    monkeypatch.setattr(get_settings(), "cache_enabled", True)
    cache.clear_local()
    try:
        before = {plan["code"] for plan in client.get("/api/v1/plans").json()}
        code = f"PYTEST_C{uuid.uuid4().hex[:6].upper()}"
        created = client.post(
            f"{ADMIN_URL}/plans",
            json={
                "code": code,
                "name": "Gói cache",
                "activity_type": "YOGA",
                "duration_days": 5,
                "price": 1000,
            },
            headers=admin["headers"],
        )
        assert created.status_code == 201
        after = {plan["code"] for plan in client.get("/api/v1/plans").json()}
        assert after == before | {code}
    finally:
        cache.clear_local()


# --------------------------------------------------------------------------
# Worker
# --------------------------------------------------------------------------


def _create_job(user_id: str, job_type: str, parameters: dict) -> str:
    from app.db.session import SessionLocal
    from app.repositories import jobs as jobs_repository

    job_id = str(uuid.uuid4())
    with SessionLocal() as session_db:
        jobs_repository.create(
            session_db,
            job_id=job_id,
            job_type=job_type,
            requested_by=user_id,
            parameters=parameters,
        )
    return job_id


def test_worker_sweep_runs_stale_pending_jobs(admin, db) -> None:
    from app.tasks import worker

    job_id = _create_job(
        admin["user"]["id"], "export.report", {"report": "products", "format": "csv"}
    )
    with db.cursor() as cursor:
        cursor.execute(
            "UPDATE background_jobs SET created_at = created_at - INTERVAL 10 MINUTE WHERE id = %s",
            (job_id,),
        )
    assert worker.sweep() >= 1
    job = client.get(f"{ADMIN_URL}/jobs/{job_id}", headers=admin["headers"]).json()
    assert job["status"] == "SUCCESS"
    assert job["progress"] == 100


def test_failed_and_unknown_jobs_are_recorded(admin) -> None:
    from app.tasks import queue

    unknown = _create_job(admin["user"]["id"], "export.nothing", {})
    queue.run(unknown)
    broken = _create_job(admin["user"]["id"], "export.report", {"report": "no_such_report"})
    queue.run(broken)
    # Tác vụ đã kết thúc thì gọi lại cũng không chạy lần nữa.
    queue.run(broken)

    first = client.get(f"{ADMIN_URL}/jobs/{unknown}", headers=admin["headers"]).json()
    assert first["status"] == "FAILED"
    assert "không được hỗ trợ" in first["error_message"]
    second = client.get(f"{ADMIN_URL}/jobs/{broken}", headers=admin["headers"]).json()
    assert second["status"] == "FAILED"
    assert "no_such_report" in second["error_message"]
